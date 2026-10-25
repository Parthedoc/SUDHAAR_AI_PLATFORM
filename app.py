from __future__ import annotations

import html
import io
import json
import os
import re
import tempfile
from pathlib import Path
from typing import Any

import streamlit as st
import streamlit.components.v1 as components
from dotenv import load_dotenv

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env", override=True)

from sudhaar.ai_engine import answer_question, load_knowledge, retrieve, ollama_status, OLLAMA_MODEL  # noqa: E402
from sudhaar.i18n import LANGUAGES, AI_RESPONSE_LANGUAGES, t  # noqa: E402
from sudhaar.safety import detect_risk, safety_copy  # noqa: E402
from sudhaar.ui_localization import install_streamlit_localization  # noqa: E402

st.set_page_config(
    page_title="SUDHAAR v0.9 — multilingual interface",
    page_icon="✦",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ----------------------------- Session state ---------------------------------
DEFAULTS: dict[str, Any] = {
    "nav_page": "home",  # retained for migration from v0.8 sessions
    "theme_mode": "Midnight",
    "ui_language_label": "English",
    "reply_language_label": "Same as my message",
    "messages": [],
    "chat_draft": "",
    "clear_draft_next_run": False,
    "upload_epoch": 0,
    "audio_epoch": 0,
    "clear_upload_next_run": False,
    "clear_audio_next_run": False,
    "feedback_done": False,
    "conversation_mode": "auto",
    "pending_transcript": "",
    "clear_transcript_next_run": False,
    "response_style": "Gentle and reassuring",
    "support_goal": "Take one manageable next step",
    "support_person": "A trusted person or counsellor",
    "support_steps": [],
    "support_plan_ready": False,
    "last_transcript": "",
}
for key, value in DEFAULTS.items():
    if key not in st.session_state:
        st.session_state[key] = value

# Keep the selected page in a non-widget state key. Streamlit may recreate a widget
# when its translated label/format changes; this canonical value prevents a language
# switch from resetting navigation to the first radio option (Home).
if "active_page" not in st.session_state:
    st.session_state["active_page"] = st.session_state.get("nav_page", "home")

# Translate app-authored UI strings using the currently selected interface language.
# Stored widget values, user messages, AI answers, source quotations and URLs remain intact.
install_streamlit_localization(st)

# Clear state before those widgets are created again; avoids Streamlit widget-state errors.
if st.session_state.pop("clear_draft_next_run", False):
    st.session_state["chat_draft"] = ""
if st.session_state.pop("clear_upload_next_run", False):
    st.session_state["upload_epoch"] += 1
if st.session_state.pop("clear_audio_next_run", False):
    st.session_state["audio_epoch"] += 1
if st.session_state.pop("clear_transcript_next_run", False):
    st.session_state["pending_transcript"] = ""
    st.session_state["transcribed_text"] = ""

UI_LANGUAGE_LABELS = list(LANGUAGES.keys())
AI_LANGUAGE_TO_NAME = {
    "Same as my message": "the same language as the user's message",
    "Assamese (অসমীয়া)": "Assamese", "Bengali (বাংলা)": "Bengali", "Bodo (बड़ो)": "Bodo",
    "Dogri (डोगरी)": "Dogri", "Gujarati (ગુજરાતી)": "Gujarati", "Hindi (हिन्दी)": "Hindi",
    "Kannada (ಕನ್ನಡ)": "Kannada", "Kashmiri (कॉशुर / كٲشُر)": "Kashmiri", "Konkani (कोंकणी)": "Konkani",
    "Maithili (मैथिली)": "Maithili", "Malayalam (മലയാളം)": "Malayalam", "Manipuri / Meitei (মৈতৈলোন্)": "Manipuri (Meitei)",
    "Marathi (मराठी)": "Marathi", "Nepali (नेपाली)": "Nepali", "Odia (ଓଡ଼ିଆ)": "Odia", "Punjabi (ਪੰਜਾਬੀ)": "Punjabi",
    "Sanskrit (संस्कृतम्)": "Sanskrit", "Santali (ᱥᱟᱱᱛᱟᱲᱤ)": "Santali", "Sindhi (سنڌي)": "Sindhi",
    "Tamil (தமிழ்)": "Tamil", "Telugu (తెలుగు)": "Telugu", "Urdu (اردو)": "Urdu", "English": "English",
}
VOICE_LOCALES = {
    "the same language as the user's message": "en-IN", "English": "en-IN", "Hindi": "hi-IN", "Bengali": "bn-IN",
    "Bodo": "hi-IN", "Dogri": "hi-IN", "Gujarati": "gu-IN", "Kannada": "kn-IN", "Kashmiri": "ks-IN",
    "Konkani": "kok-IN", "Maithili": "mai-IN", "Malayalam": "ml-IN", "Manipuri (Meitei)": "mni-IN",
    "Marathi": "mr-IN", "Nepali": "ne-IN", "Odia": "or-IN", "Punjabi": "pa-IN", "Sanskrit": "sa-IN",
    "Santali": "sat-IN", "Sindhi": "sd-IN", "Tamil": "ta-IN", "Telugu": "te-IN", "Urdu": "ur-IN", "Assamese": "as-IN",
}

NAV_LABELS = {
    "home": ("⌂", "Home", "A calm starting point"),
    "chat": ("✦", "SUDHAAR Companion", "Talk without judgement"),
    "library": ("◈", "Health Library", "Understand the evidence"),
    "services": ("♡", "Find Support", "People and services"),
    "programmes": ("↗", "Programmes & Apps", "Official pathways"),
    "plan": ("☼", "My Next Step", "Your plan, your pace"),
    "settings": ("⚙", "Preferences", "Make it yours"),
    "feedback": ("✎", "Feedback", "Help us improve"),
}

PALETTES = {
    "Midnight": {
        "bg": "#070d18", "bg2": "#0d1728", "panel": "#101c30", "panel2": "#162640",
        "sidebar": "#091120", "input": "#0e192b", "ink": "#f1f5fc", "body": "#c7d3e5",
        "muted": "#93a4bf", "line": "rgba(143,166,204,.20)", "accent": "#56ddbd", "accent2": "#91adff",
        "hero1": "rgba(37,70,119,.68)", "hero2": "rgba(19,72,76,.50)", "warning": "#332b18",
        "shadow": "rgba(0,0,0,.24)", "onaccent": "#071b1b", "soft": "rgba(86,221,189,.09)",
    },
    "Pearl": {
        "bg": "#f4f7fc", "bg2": "#edf2f9", "panel": "#ffffff", "panel2": "#f3f7fd",
        "sidebar": "#ffffff", "input": "#ffffff", "ink": "#17253b", "body": "#344761",
        "muted": "#667892", "line": "rgba(36,57,89,.16)", "accent": "#087c68", "accent2": "#365bb4",
        "hero1": "rgba(217,231,255,.95)", "hero2": "rgba(211,245,237,.90)", "warning": "#fff3ce",
        "shadow": "rgba(39,61,93,.10)", "onaccent": "#ffffff", "soft": "rgba(8,124,104,.08)",
    },
}


def build_css(theme: str) -> str:
    p = PALETTES.get(theme, PALETTES["Midnight"])
    vars_css = ";".join(f"--{key}:{value}" for key, value in p.items())
    return f"""<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Manrope:wght@500;600;700;800&display=swap');
:root{{{vars_css}}}
html,body,[class*="css"]{{font-family:'DM Sans',system-ui,sans-serif;scroll-behavior:smooth;}}
.stApp{{background:radial-gradient(ellipse at 2% 0%,var(--hero1) 0%,transparent 34%),radial-gradient(ellipse at 100% 16%,var(--hero2) 0%,transparent 28%),linear-gradient(155deg,var(--bg2),var(--bg) 65%);color:var(--ink);}}
/* Streamlit chrome: hide Deploy, but keep the header that owns sidebar controls. */
[data-testid="stAppDeployButton"], #MainMenu, footer{{display:none!important;}}
[data-testid="stToolbar"], [data-testid="stHeaderActionElements"]{{visibility:visible!important;opacity:1!important;pointer-events:auto!important;}}
/* Keep the sidebar collapse/reopen control visible and clickable in both states. */
[data-testid="stSidebarCollapseButton"], [data-testid="stSidebarCollapsedControl"], [data-testid="collapsedControl"]{{visibility:visible!important;opacity:1!important;pointer-events:auto!important;z-index:999999!important;}}
[data-testid="stSidebarCollapseButton"]{{display:flex!important;}}
[data-testid="stSidebarCollapsedControl"], [data-testid="collapsedControl"]{{display:flex!important;position:fixed!important;top:.55rem!important;left:.55rem!important;}}
[data-testid="stSidebarCollapseButton"] button, [data-testid="stSidebarCollapsedControl"] button, [data-testid="collapsedControl"] button{{color:var(--ink)!important;background:var(--panel)!important;border:1px solid var(--line)!important;border-radius:10px!important;min-width:34px!important;min-height:34px!important;box-shadow:0 5px 16px var(--shadow)!important;}}
header[data-testid="stHeader"]{{min-height:2.6rem;height:2.6rem;background:transparent!important;z-index:1000;}}
[data-testid="stSidebar"]{{background:var(--sidebar);border-right:1px solid var(--line);}}
[data-testid="stSidebar"]>div:first-child{{padding-top:1.25rem;}}
.block-container{{max-width:2200px;width:98.5%;padding-top:.85rem;padding-bottom:6rem;}}
.stApp, .stApp p, .stApp li, .stApp label, .stApp [data-testid="stMarkdownContainer"]{{color:var(--body);}}
h1,h2,h3,h4{{font-family:'Manrope',sans-serif;letter-spacing:-.035em;color:var(--ink)!important;}}
h1{{font-size:clamp(2rem,3.6vw,3.35rem)!important;font-weight:800!important;}}
h2{{font-size:clamp(1.4rem,2.2vw,2rem)!important;font-weight:800!important;}}
p,li{{line-height:1.72;}}
strong,b{{color:var(--ink);}}
a{{color:var(--accent)!important;text-decoration:none!important;}}
a:hover{{text-decoration:underline!important;}}
.stCaption,small,.muted{{color:var(--muted)!important;}}
[data-testid="stAlert"]{{border-radius:15px;border:1px solid var(--line);}}
[data-testid="stWarning"]{{background:var(--warning);}}
.stButton>button,.stLinkButton>a{{border-radius:13px!important;min-height:42px;border:1px solid var(--line)!important;background:var(--panel2)!important;color:var(--ink)!important;font-weight:700;transition:transform .22s ease,box-shadow .22s ease,border-color .22s ease,background .22s ease;}}
.stButton>button:hover,.stLinkButton>a:hover{{transform:translateY(-2px);border-color:var(--accent)!important;box-shadow:0 9px 25px var(--shadow);}}
.stButton>button:focus-visible,.stLinkButton>a:focus-visible{{outline:3px solid var(--accent2)!important;outline-offset:2px;}}
.stTextInput input,.stTextArea textarea,[data-baseweb="select"]>div{{background:var(--input)!important;color:var(--ink)!important;border:1px solid var(--line)!important;border-radius:12px!important;}}
.stTextInput input::placeholder,.stTextArea textarea::placeholder{{color:var(--muted)!important;opacity:.9;}}
[data-testid="stChatMessage"]{{background:transparent;border:0;border-radius:0;margin:.7rem 0 1.1rem;box-shadow:none;animation:messageIn .4s ease both;width:100%;max-width:none;}}
[data-testid="stChatMessageContent"]{{max-width:100%!important;width:100%;}}
[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"]{{max-width:100%;}}
.user-bubble-v6{{display:inline-block;max-width:min(900px,95%);margin:0 0 0 auto;padding:13px 17px;border:1px solid var(--line);border-radius:20px 20px 5px 20px;background:var(--panel2);color:var(--ink);line-height:1.7;white-space:pre-wrap;overflow-wrap:anywhere;}}
.user-bubble-row-v6{{display:flex;justify-content:flex-end;width:100%;margin:.9rem 0 1.1rem;}}
.assistant-heading-v6{{display:flex;align-items:center;gap:9px;color:var(--accent);font-size:.78rem;font-weight:800;letter-spacing:.08em;margin:0 0 .35rem;}}
.assistant-mark-v6{{width:30px;height:30px;display:inline-flex;align-items:center;justify-content:center;border-radius:10px;background:linear-gradient(135deg,var(--accent),var(--accent2));color:var(--onaccent);font-family:Manrope;font-weight:900;letter-spacing:0;}}
.chat-intro-v6{{max-width:1100px;margin:0 auto 1.35rem;text-align:center;padding:clamp(22px,3vw,38px);border:1px solid var(--line);border-radius:28px;background:linear-gradient(125deg,var(--panel),var(--panel2));box-shadow:0 15px 45px var(--shadow);}}
.chat-intro-v6 .sub{{color:var(--muted);font-size:.98rem;line-height:1.7;}}
.chat-bottom-tools-v6{{border:1px solid var(--line);border-radius:18px;background:var(--panel);padding:10px 13px;margin:.4rem 0 .7rem;}}
.chat-tools-heading{{display:flex;align-items:center;gap:8px;margin:1.5rem 0 .65rem;font-size:.88rem;font-weight:800;color:var(--ink);}}
.chat-tools-heading span:first-child{{display:inline-flex;align-items:center;justify-content:center;background:var(--soft);color:var(--accent);width:28px;height:28px;border-radius:9px;font-size:1.1rem;}}
[data-testid="stTabs"] [data-baseweb="tab-list"]{{gap:8px;border-bottom:1px solid var(--line);}}
[data-testid="stTabs"] button[role="tab"]{{color:var(--body)!important;}}
[data-testid="stTabs"] button[role="tab"]{{border-radius:12px 12px 0 0;padding:10px 16px;font-weight:700;}}
[data-testid="stTabs"] button[role="tab"][aria-selected="true"]{{background:var(--soft);color:var(--accent)!important;}}
.assistant-answer-body-v7{{width:100%;max-width:none;font-size:1.02rem;line-height:1.82;padding:0 0 .25rem;}}
.assistant-answer-body-v7 p,.assistant-answer-body-v7 li{{max-width:none;}}
.stApp [data-testid="stMarkdownContainer"]{{max-width:100%!important;}}
.stApp [data-testid="stVerticalBlock"]{{gap: .6rem;}}
[data-testid="stChatInput"]{{width:100%;}}
[data-testid="stChatInput"] textarea{{background:var(--input)!important;color:var(--ink)!important;border:1px solid var(--line)!important;border-radius:18px!important;min-height:58px!important;}}
.stChatInputContainer{{max-width:none!important;padding-bottom:.4rem;}}
.stChatInputContainer textarea{{border-radius:18px!important;min-height:56px!important;background:var(--input)!important;color:var(--ink)!important;}}
[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] p{{color:var(--body);}}
[data-testid="stExpander"]{{border:1px solid var(--line);border-radius:14px;background:var(--panel);}}
/* Keep labels and popovers readable in both theme modes. */
[data-testid="stWidgetLabel"], [data-testid="stWidgetLabel"] *, .stRadio label, .stCheckbox label, .stSelectbox label, .stMultiSelect label{{color:var(--ink)!important;}}
[data-baseweb="popover"], [data-baseweb="menu"], ul[role="listbox"], li[role="option"]{{background:var(--panel)!important;color:var(--ink)!important;}}
li[role="option"] *, [role="option"]{{color:var(--ink)!important;}}
[data-baseweb="select"] *, [data-baseweb="input"] *, [data-baseweb="textarea"] *{{color:var(--ink)!important;}}
.stApp [data-testid="stMarkdownContainer"] p, .stApp [data-testid="stMarkdownContainer"] li{{color:var(--body);}}
.stApp [data-testid="stMarkdownContainer"] strong, .stApp [data-testid="stMarkdownContainer"] b{{color:var(--ink)!important;}}
@media (max-width: 700px){{.block-container{{width:100%;padding-left:.7rem;padding-right:.7rem;}}.chat-intro{{padding:16px;}}.hero{{padding:22px;border-radius:20px;}}.user-bubble-v6{{max-width:100%;}}}}
[data-testid="stFileUploaderDropzone"]{{background:var(--panel2);border:1px dashed var(--accent2);border-radius:16px;}}
[data-testid="stAudioInput"]{{background:var(--panel2);border:1px solid var(--line);border-radius:15px;padding:.35rem;}}
hr{{border-color:var(--line)!important;margin:1.35rem 0!important;}}
/* Brand and high-touch UI components */
.brand-row{{display:flex;align-items:center;gap:11px;margin:0 0 18px;padding:0 2px;}}
.brand-mark{{width:46px;height:46px;display:flex;align-items:center;justify-content:center;border-radius:16px;background:linear-gradient(135deg,var(--accent),var(--accent2));color:var(--onaccent);font-family:'Manrope';font-weight:900;font-size:23px;box-shadow:0 8px 26px var(--shadow);}}
.brand-name{{font-family:'Manrope';font-size:1.45rem;letter-spacing:-.045em;font-weight:800;color:var(--ink);line-height:1.08;}}
.brand-sub{{font-size:.63rem;letter-spacing:.15em;color:var(--muted);font-weight:800;margin-top:5px;}}
.side-overline{{font-size:.65rem;font-weight:800;letter-spacing:.17em;text-transform:uppercase;color:var(--muted);margin:.75rem .2rem .45rem;}}
.side-status{{display:flex;gap:9px;align-items:center;border:1px solid var(--line);border-radius:14px;padding:11px 12px;background:var(--panel);font-size:.78rem;color:var(--body);margin:.85rem 0;}}
.status-dot{{width:8px;height:8px;border-radius:50%;background:var(--accent);box-shadow:0 0 0 4px var(--soft);flex:0 0 auto;}}
.hero{{position:relative;overflow:hidden;border:1px solid var(--line);border-radius:30px;padding:clamp(25px,3.6vw,48px);background:linear-gradient(124deg,var(--hero1),var(--panel) 58%,var(--hero2));box-shadow:0 25px 85px var(--shadow);margin:.25rem 0 25px;animation:heroIn .72s cubic-bezier(.2,.75,.25,1) both;}}
.hero:after{{content:"";position:absolute;right:-96px;top:-170px;width:390px;height:390px;border-radius:50%;border:1px solid var(--line);box-shadow:0 0 0 34px var(--soft),0 0 0 80px var(--soft);pointer-events:none;}}
.eyebrow,.section-kicker{{font-size:.68rem;letter-spacing:.17em;text-transform:uppercase;font-weight:800;color:var(--accent);}}
.hero h1{{font-size:clamp(3rem,6vw,5rem)!important;line-height:.98;margin:.8rem 0 1.1rem;}}
.hero p{{color:var(--body);font-size:1.02rem;max-width:790px;line-height:1.85;}}
.pill{{display:inline-flex;align-items:center;gap:5px;padding:6px 11px;border:1px solid var(--line);border-radius:999px;background:var(--soft);font-size:.72rem;font-weight:700;color:var(--accent);margin:3px 5px 3px 0;}}
.note-label{{margin-top:15px;font-size:.68rem;letter-spacing:.1em;color:var(--muted);font-weight:800;}}
.card-grid{{display:grid;gap:14px;}}
.metric-card,.feature-card,.resource-card,.kb-card,.programme-card,.plan-card,.companion-card{{border:1px solid var(--line);border-radius:20px;padding:19px;background:var(--panel);height:100%;box-shadow:0 9px 26px var(--shadow);transition:transform .25s ease,border-color .25s ease,box-shadow .25s ease;}}
.metric-card:hover,.feature-card:hover,.resource-card:hover,.kb-card:hover,.programme-card:hover,.plan-card:hover,.companion-card:hover{{transform:translateY(-4px);border-color:var(--accent);box-shadow:0 16px 40px var(--shadow);}}
.metric-num{{font-family:'Manrope';font-weight:800;font-size:1.9rem;color:var(--accent);letter-spacing:-.04em;}}
.metric-label{{font-size:.82rem;color:var(--muted);margin-top:4px;}}
.icon-box{{width:43px;height:43px;display:flex;align-items:center;justify-content:center;border-radius:15px;background:var(--soft);color:var(--accent);font-size:1.15rem;margin-bottom:12px;}}
.card-title{{font-family:'Manrope';font-size:1.02rem;font-weight:800;color:var(--ink);margin-bottom:7px;}}
.card-text{{font-size:.9rem;line-height:1.68;color:var(--body);}}
.muted{{color:var(--muted);font-size:.79rem;}}
.section-head{{display:flex;align-items:flex-end;justify-content:space-between;gap:16px;margin:1.5rem 0 1rem;}}
.section-head h2{{margin:0;}}
.chat-intro{{border:1px solid var(--line);border-radius:23px;padding:21px 23px;background:linear-gradient(120deg,var(--panel),var(--panel2));margin:.3rem 0 1rem;}}
.chat-intro strong{{font-size:1.07rem;color:var(--ink);}}
.chat-intro p{{margin:.3rem 0 0;color:var(--body);}}
.mode-card{{border:1px solid var(--line);border-radius:16px;background:var(--panel2);padding:12px 14px;height:100%;}}
.mode-card .mode-icon{{font-size:1.25rem;color:var(--accent);}}
.mode-card .mode-title{{font-weight:800;color:var(--ink);margin:.3rem 0;}}
.mode-card .mode-desc{{font-size:.8rem;color:var(--muted);line-height:1.55;}}
.chat-compose{{border:1px solid var(--line);background:var(--panel);padding:15px;border-radius:22px;box-shadow:0 14px 45px var(--shadow);}}
.source-chip{{display:inline-block;margin:3px 5px 3px 0;padding:7px 10px;border:1px solid var(--line);background:var(--panel2);border-radius:10px;font-size:.77rem;}}
.anim-in{{animation:riseIn .6s cubic-bezier(.2,.7,.2,1) both;}}
@keyframes heroIn{{from{{opacity:0;transform:translateY(16px) scale(.99)}}to{{opacity:1;transform:translateY(0) scale(1)}}}}
@keyframes riseIn{{from{{opacity:0;transform:translateY(14px)}}to{{opacity:1;transform:translateY(0)}}}}
@keyframes messageIn{{from{{opacity:0;transform:translateY(8px)}}to{{opacity:1;transform:translateY(0)}}}}
@supports (animation-timeline:view()){{.scroll-reveal{{animation:scrollReveal linear both;animation-timeline:view();animation-range:entry 0% entry 42%;}}@keyframes scrollReveal{{from{{opacity:.35;transform:translateY(19px)}}to{{opacity:1;transform:translateY(0)}}}}}}
@media(max-width:850px){{.block-container{{padding-left:1rem;padding-right:1rem;}}.hero{{padding:24px;border-radius:22px;}}.hero:after{{opacity:.3;}}}}
@media(prefers-reduced-motion:reduce){{*,*::before,*::after{{animation:none!important;scroll-behavior:auto!important;transition:none!important;}}}}
</style>"""


def safe_html_text(value: str) -> str:
    return html.escape(str(value), quote=True)


def sync_active_page_from_radio() -> None:
    """Persist the page selected in the translated navigation widget."""
    selected = st.session_state.get("nav_radio_widget", st.session_state.get("active_page", "home"))
    if selected in NAV_LABELS:
        st.session_state["active_page"] = selected
        st.session_state["nav_page"] = selected  # compatibility with older code/session state


def preserve_active_page_on_language_change() -> None:
    """Capture the current page before a locale change can recreate radio widgets."""
    selected = st.session_state.get("nav_radio_widget", st.session_state.get("active_page", "home"))
    if selected in NAV_LABELS:
        st.session_state["active_page"] = selected
        st.session_state["nav_page"] = selected


def go_to_page(target: str) -> None:
    # Canonical page state is independent of translated widget labels.
    if target in NAV_LABELS:
        st.session_state["active_page"] = target
        st.session_state["nav_page"] = target
        st.session_state["nav_radio_widget"] = target


def clear_chat_and_go_to_chat() -> None:
    st.session_state["messages"] = []
    st.session_state["clear_draft_next_run"] = True
    st.session_state["clear_upload_next_run"] = True
    st.session_state["clear_audio_next_run"] = True
    st.session_state["active_page"] = "chat"
    st.session_state["nav_page"] = "chat"
    st.session_state["nav_radio_widget"] = "chat"


def tts_button(text: str, voice_locale: str, widget_key: str) -> None:
    """Browser speech synthesis. Available voices depend on Windows/browser language packs."""
    safe_text = json.dumps(text, ensure_ascii=False).replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
    safe_locale = json.dumps(voice_locale)
    html_doc = f"""<!doctype html><html><head><meta charset='utf-8'><style>
    body{{margin:0;font-family:Arial,sans-serif;background:transparent}}
    button{{border:1px solid rgba(120,150,190,.45);background:rgba(100,150,170,.10);color:inherit;border-radius:10px;padding:8px 12px;font-size:12px;font-weight:700;cursor:pointer;margin-right:6px}}
    button:hover{{border-color:#41bda3;background:rgba(65,189,163,.15)}}
    </style></head><body>
    <button id='speak'>🔊 Listen to response</button><button id='stop'>■ Stop</button>
    <script>
    const txt={safe_text}; const lang={safe_locale};
    document.getElementById('speak').addEventListener('click',()=>{{
      if(!('speechSynthesis' in window)){{alert('Text-to-speech is not supported by this browser.');return;}}
      window.speechSynthesis.cancel(); const u=new SpeechSynthesisUtterance(txt); u.lang=lang; u.rate=0.94; u.pitch=1.0;
      const voices=window.speechSynthesis.getVoices(); const prefix=lang.split('-')[0].toLowerCase();
      const v=voices.find(x=>x.lang.toLowerCase()===lang.toLowerCase())||voices.find(x=>x.lang.toLowerCase().startsWith(prefix));
      if(v)u.voice=v; window.speechSynthesis.speak(u);
    }});
    document.getElementById('stop').addEventListener('click',()=>window.speechSynthesis.cancel());
    </script></body></html>"""
    components.html(html_doc, height=47, scrolling=False, tab_index=0)


def extract_uploaded_files(uploaded_files) -> tuple[str, list[str], list[str]]:
    """Read text from user-uploaded files locally; returns context, names, warnings."""
    if not uploaded_files:
        return "", [], []
    excerpts: list[str] = []
    names: list[str] = []
    warnings: list[str] = []
    total_chars = 0
    max_size = 8 * 1024 * 1024
    for item in uploaded_files:
        name = Path(item.name).name
        names.append(name)
        if getattr(item, "size", 0) > max_size:
            warnings.append(f"{name}: skipped because it is larger than 8 MB.")
            continue
        suffix = Path(name).suffix.casefold()
        try:
            raw = item.getvalue()
            if suffix in {".txt", ".md"}:
                text = raw.decode("utf-8", errors="replace")
            elif suffix == ".pdf":
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(raw))
                pieces = []
                for page in reader.pages[:20]:
                    pieces.append(page.extract_text() or "")
                text = "\n".join(pieces)
            elif suffix == ".docx":
                from docx import Document
                document = Document(io.BytesIO(raw))
                text = "\n".join(p.text for p in document.paragraphs)
                for table in document.tables[:10]:
                    for row in table.rows:
                        text += "\n" + " | ".join(cell.text for cell in row.cells)
            else:
                warnings.append(f"{name}: unsupported format.")
                continue
            text = text.strip()
            if not text:
                warnings.append(f"{name}: no extractable text found (scanned PDFs need OCR, which this prototype does not do).")
                continue
            allowance = max(0, 9500 - total_chars)
            if allowance <= 0:
                warnings.append("Only the first 9,500 characters across uploaded files are included per message.")
                break
            text = text[:allowance]
            total_chars += len(text)
            excerpts.append(f"[User-uploaded document: {name} — NOT independently verified]\n{text}")
        except Exception as exc:
            warnings.append(f"{name}: could not extract text ({type(exc).__name__}).")
    return "\n\n".join(excerpts), names, warnings


@st.cache_data(ttl=10, show_spinner=False)
def cached_ollama_status():
    return ollama_status()


@st.cache_resource(show_spinner=False)
def get_whisper_model(model_size: str = "base"):
    from faster_whisper import WhisperModel
    # CPU int8 is a pragmatic low-cost Windows default; multilingual transcription model is downloaded on first use.
    return WhisperModel(model_size, device="cpu", compute_type="int8")


def transcribe_audio(audio_bytes: bytes, suffix: str = ".wav") -> tuple[str, str | None]:
    """Transcribe recorded or uploaded audio locally using optional faster-whisper."""
    temp_path = None
    try:
        model_size = os.getenv("SUDHAAR_WHISPER_MODEL", "base")
        model = get_whisper_model(model_size)
        safe_suffix = suffix if suffix and re.fullmatch(r"\.[A-Za-z0-9]{2,5}", suffix) else ".wav"
        with tempfile.NamedTemporaryFile(delete=False, suffix=safe_suffix) as temp:
            temp.write(audio_bytes)
            temp_path = temp.name
        segments, info = model.transcribe(temp_path, language=None, task="transcribe", vad_filter=True, beam_size=3)
        text = " ".join(segment.text.strip() for segment in segments).strip()
        return text, getattr(info, "language", None)
    finally:
        if temp_path:
            try:
                os.unlink(temp_path)
            except OSError:
                pass


def infer_voice_locale(text: str, selected_reply_language: str) -> str:
    """Choose an installed browser speech locale based on the chosen language or writing system."""
    if selected_reply_language != "the same language as the user's message":
        return VOICE_LOCALES.get(selected_reply_language, "en-IN")
    text = text or ""
    script_checks = [
        (r"[\u0900-\u097f]", "hi-IN"), (r"[\u0980-\u09ff]", "bn-IN"),
        (r"[\u0a80-\u0aff]", "gu-IN"), (r"[\u0a00-\u0a7f]", "pa-IN"),
        (r"[\u0b00-\u0b7f]", "or-IN"), (r"[\u0b80-\u0bff]", "ta-IN"),
        (r"[\u0c00-\u0c7f]", "te-IN"), (r"[\u0c80-\u0cff]", "kn-IN"),
        (r"[\u0d00-\u0d7f]", "ml-IN"), (r"[\u0600-\u06ff]", "ur-IN"),
        (r"[\u0d80-\u0dff]", "si-LK"),
    ]
    import re
    for pattern, locale in script_checks:
        if re.search(pattern, text):
            return locale
    return "en-IN"


def display_sources(sources: list[dict[str, Any]], language: str) -> None:
    if not sources:
        return
    with st.expander(f"◈  Evidence & sources ({len(sources)})", expanded=False):
        st.caption(t("source_note", language))
        for source in sources:
            st.markdown(f"**{source.get('title', 'Source')}**  ")
            st.caption(f"{source.get('source_name', 'Source not listed')} · {source.get('jurisdiction', 'Jurisdiction not listed')}")
            if source.get("content"):
                st.write(source["content"])
            if source.get("source_url"):
                st.markdown(f"[Open original source ↗]({source['source_url']})")
            st.markdown("---")


def process_user_message(prompt: str, response_language: str, language_code: str, conversation_mode: str, response_style: str, extra_context: str = "") -> None:
    prompt = (prompt or "").strip()
    if not prompt:
        st.warning("Write a message first, or use the microphone to transcribe one.")
        return
    prior = st.session_state.messages[-8:]
    st.session_state.messages.append({"role": "user", "content": prompt, "kind": "user"})
    risk = detect_risk(prompt)
    if risk:
        copy = safety_copy(risk, language_code)
        answer = f"**{copy['title']}**\n\n{copy['body']}\n\n{copy['main']}\n\n"
        for contact in copy.get("contacts", []):
            name = contact.get("name", "Support service")
            url = contact.get("url", "https://112.gov.in/")
            answer += f"- **{name}** — [official information]({url})\n"
        sources = retrieve(prompt)
        mode = "safety"
        error = ""
    else:
        with st.spinner("SUDHAAR is here with you · finding relevant sources and preparing a thoughtful response…"):
            answer, sources, mode, error = answer_question(
                prompt, lang=language_code, history=prior, response_language=response_language,
                conversation_mode=conversation_mode, extra_context=extra_context, style=response_style,
            )
    st.session_state.messages.append({"role": "assistant", "content": answer, "sources": sources, "mode": mode, "error": error})
    st.session_state["clear_draft_next_run"] = True
    st.session_state["clear_upload_next_run"] = True
    st.session_state["clear_audio_next_run"] = True
    st.session_state["clear_transcript_next_run"] = True
    st.rerun()


# ------------------------------- Sidebar --------------------------------------
with st.sidebar:
    st.markdown(
        '<div class="brand-row"><div class="brand-mark">S</div><div><div class="brand-name">SUDHAAR</div>'
        '<div class="brand-sub">A SPACE WITHOUT JUDGEMENT</div></div></div>',
        unsafe_allow_html=True,
    )
    selected_ui = st.selectbox(
        "Interface language / इंटरफ़ेस भाषा", UI_LANGUAGE_LABELS,
        key="ui_language_label", on_change=preserve_active_page_on_language_change,
    )
    lang = LANGUAGES[selected_ui]
    st.caption(t("brand_tag", lang))
    st.selectbox("AI reply language", AI_RESPONSE_LANGUAGES, key="reply_language_label", help="This asks the local model to respond in the chosen language. Quality varies and must be reviewed, especially for medical and crisis wording.")
    st.selectbox("Appearance", ["Midnight", "Pearl"], key="theme_mode", format_func=lambda value: "🌙 Midnight · dark" if value == "Midnight" else "☀️ Pearl · light")
    st.markdown('<div class="side-overline">YOUR SPACE</div>', unsafe_allow_html=True)
    nav_ids = list(NAV_LABELS.keys())
    # Re-sync the widget key from canonical state BEFORE creating the widget. The
    # visible label changes by language; the selected page ID must not.
    active_page = st.session_state.get("active_page", "home")
    if active_page not in nav_ids:
        active_page = "home"
        st.session_state["active_page"] = active_page
    st.session_state["nav_radio_widget"] = active_page
    page = st.radio(
        "Your space", nav_ids,
        format_func=lambda item: f"{NAV_LABELS[item][0]}  {NAV_LABELS[item][1]}",
        key="nav_radio_widget", on_change=sync_active_page_from_radio,
        label_visibility="collapsed",
    )
    st.session_state["active_page"] = page
    st.session_state["nav_page"] = page
    local_status = cached_ollama_status()
    if local_status["server"] and local_status["model"]:
        status_text = f"Local AI ready · {OLLAMA_MODEL}"
        status_color = "ready"
    elif local_status["server"]:
        status_text = f"Ollama is running · model missing"
        status_color = "warn"
    else:
        status_text = "Local AI offline · start Ollama"
        status_color = "warn"
    st.markdown(f'<div class="side-status"><span class="status-dot"></span><span>{safe_html_text(status_text)}</span></div>', unsafe_allow_html=True)
    st.caption("Your chat stays in this running session; the app does not save it to a database. If you upload documents, their extracted text is sent to your local Ollama model only.")

st.markdown(build_css(st.session_state.theme_mode), unsafe_allow_html=True)

# Urdu interface layout uses right-to-left reading order; AI reply language remains a separate preference.
if lang == "ur":
    st.markdown(
        "<style>.block-container{direction:rtl;text-align:right;} "
        "[data-testid=\"stSidebar\"]{direction:rtl;text-align:right;} "
        "[data-testid=\"stWidgetLabel\"],.stMarkdownContainer{text-align:right;} "
        "</style>",
        unsafe_allow_html=True,
    )

# Emergency support is visible but intentionally calm, non-alarmist.
st.markdown(
    f'<div style="border:1px solid var(--line);background:var(--warning);border-radius:14px;padding:10px 14px;display:flex;gap:10px;align-items:center;margin-bottom:16px;">'
    f'<span style="font-size:1.1rem">♡</span><span style="font-size:.84rem;color:var(--ink)">{safe_html_text(t("safety_banner", lang))}</span></div>',
    unsafe_allow_html=True,
)

reply_lang = AI_LANGUAGE_TO_NAME.get(st.session_state.reply_language_label, "the same language as the user's message")
voice_locale = VOICE_LOCALES.get(reply_lang, "en-IN")
kb = load_knowledge()
resources_path = BASE_DIR / "sudhaar" / "data" / "resources.json"
try:
    resources = json.loads(resources_path.read_text(encoding="utf-8"))
except (OSError, json.JSONDecodeError):
    resources = []

# --------------------------------- Pages --------------------------------------
if page == "home":
    st.markdown(
        f'<div class="hero"><div class="eyebrow">SUBSTANCE USE DE-ADDICTION · HELP · AWARENESS · RECOVERY</div>'
        f'<h1>SUDHAAR<span style="color:var(--accent)">.</span></h1>'
        f'<p><b>A softer place to start.</b><br>Whatever brought you here, you deserve information and support without shame. Talk things through, learn about health, or find a real-world next step — at your own pace.</p>'
        f'<span class="pill">✦ No judgement</span><span class="pill">◈ Evidence-linked</span><span class="pill">♡ Your pace</span><span class="pill">⌂ Local-first</span>'
        f'<div class="note-label">A LOCAL AI PROTOTYPE · NOT A CLINICAL OR EMERGENCY SERVICE</div></div>',
        unsafe_allow_html=True,
    )
    c1, c2 = st.columns([1.3, 1])
    with c1:
        st.markdown('<div class="section-head"><div><div class="section-kicker">START WHERE YOU ARE</div><h2>What would feel helpful right now?</h2></div></div>', unsafe_allow_html=True)
        home_options = [
            ("chat", "♡", "I just need to talk", "A calm, non-judgmental place to put your thoughts into words."),
            ("library", "◈", "I want to understand something", "Clear explanations of health, treatment and recovery topics."),
            ("services", "↗", "I need help finding support", "Official helplines, programmes and care pathways."),
        ]
        for target, icon, title, desc in home_options:
            st.markdown(f'<div class="feature-card scroll-reveal" style="display:flex;gap:14px;align-items:flex-start;margin-bottom:10px"><div class="icon-box" style="flex-shrink:0;margin:0">{icon}</div><div><div class="card-title">{title}</div><div class="card-text">{desc}</div></div></div>', unsafe_allow_html=True)
            st.button(f"Open · {title}", key=f"home_{target}", use_container_width=True, on_click=go_to_page, args=(target,))
    with c2:
        st.markdown('<div class="companion-card scroll-reveal"><div class="icon-box">✦</div><div class="section-kicker">A NOTE FROM SUDHAAR</div><h2>You do not have to explain it perfectly.</h2><p>Start with one sentence, a question, or even “I don’t know where to begin.” I’ll do my best to respond kindly and help you look at one next step.</p><p class="muted">I’m an AI, not a person or a clinician. You stay in control of what you share.</p></div>', unsafe_allow_html=True)
        st.markdown('<br><div class="companion-card"><div class="section-kicker">KEEP HUMAN HELP CLOSE</div><div class="call-number">112</div><p>Immediate danger or a life-threatening emergency.</p><hr><div class="call-number">14416</div><p>Mental-health support via Tele-MANAS.</p><hr><div class="call-number">14446</div><p>National De-addiction Helpline.</p></div>', unsafe_allow_html=True)
    st.markdown('<div class="section-head"><div><div class="section-kicker">HOW IT WORKS</div><h2>Warmth, backed by sources.</h2></div></div>', unsafe_allow_html=True)
    cols = st.columns(3)
    for col, icon, title, text in zip(cols, ["♡", "◈", "↗"], ["Start with your words", "Explore reliable information", "Choose your next step"], ["Talk, ask, or upload a document you want help understanding.", "SUDHAAR retrieves relevant passages and links you back to their sources.", "Use a support plan or open official services when you are ready."]):
        with col:
            st.markdown(f'<div class="feature-card scroll-reveal"><div class="icon-box">{icon}</div><div class="card-title">{title}</div><div class="card-text">{text}</div></div>', unsafe_allow_html=True)

elif page == "chat":
    # SUDHAAR Companion: factual Q&A, emotional support, document context and voice input.
    if not st.session_state.messages:
        st.markdown(
            '<div class="chat-intro-v6 anim-in">'
            '<div class="brand-mark" style="margin:0 auto 14px">S</div>'
            '<div class="section-kicker">A SPACE WITHOUT JUDGEMENT</div>'
            '<h1 style="margin:.45rem 0 1rem">Ask anything. Start anywhere.</h1>'
            '<p class="sub">Learn about alcohol or other substances, understand health effects, explore support options, ask about a document, or talk through something difficult. SUDHAAR will try to answer the question you actually asked.</p>'
            '<span class="pill">Direct answers</span><span class="pill">Sources when available</span><span class="pill">Local AI · no API fee</span>'
            '</div>',
            unsafe_allow_html=True,
        )
        prompt_cards = [
            ("Alcohol: general knowledge", "What is alcohol, how does it affect the body, and what are its short- and long-term health risks?"),
            ("Cannabis: effects and risks", "What is cannabis and what are its possible short-term and long-term health effects?"),
            ("Understand substance dependence", "What is substance dependence, and what types of treatment and support can help?"),
            ("Help for someone I care about", "How can I support someone affected by substance use without judging or pressuring them?"),
        ]
        pcols = st.columns(2)
        for idx, (label, prompt_text) in enumerate(prompt_cards):
            with pcols[idx % 2]:
                if st.button(label + "  ↗", key=f"starter_v7_{idx}", use_container_width=True):
                    st.session_state["queued_prompt_v6"] = prompt_text
                    st.rerun()
    else:
        st.markdown('<div class="section-kicker">YOUR CONVERSATION · FULL-WIDTH ANSWERS</div>', unsafe_allow_html=True)
        tools_row = st.columns([5, 1])
        with tools_row[0]:
            st.markdown("### Talk to SUDHAAR")
            st.caption("Ask for general knowledge, upload material to discuss, or speak your message. Replies use your local Ollama model.")
        with tools_row[1]:
            if st.button("＋ New chat", key="new_chat_v7", use_container_width=True):
                st.session_state["messages"] = []
                st.session_state["clear_draft_next_run"] = True
                st.session_state["clear_upload_next_run"] = True
                st.session_state["clear_audio_next_run"] = True
                st.session_state["clear_transcript_next_run"] = True
                st.rerun()

    mode_options = {
        "auto": "✦  Answer my question / choose automatically",
        "listen": "♡  I need someone to listen",
        "health": "◈  Explain a health concern",
        "recovery": "☼  Think through a recovery step",
        "family": "⌂  Help me support someone",
        "services": "↗  Find a service",
    }
    mode_col, style_col = st.columns([1.5, 1])
    with mode_col:
        st.selectbox("How should SUDHAAR respond?", list(mode_options.keys()), format_func=lambda key: mode_options[key], key="conversation_mode")
    with style_col:
        st.selectbox("Answer detail", ["Clear and concise", "Gentle and reassuring", "Detailed explanation"], key="response_style", index=0)

    # Render each turn once. User messages are right-aligned; responses have the whole content width.
    for idx, item in enumerate(st.session_state.messages):
        if item.get("role") == "user":
            safe_text = safe_html_text(item.get("content", "")).replace("\n", "<br>")
            st.markdown(
                f'<div class="user-bubble-row-v6"><div class="user-bubble-v6">{safe_text}</div></div>',
                unsafe_allow_html=True,
            )
        else:
            st.markdown('<div class="assistant-heading-v6"><span class="assistant-mark-v6">S</span><span>SUDHAAR</span></div>', unsafe_allow_html=True)
            # Do not render the answer inside a narrow chat-message widget; normal markdown spans the page width.
            st.markdown(item.get("content", ""))
            if item.get("mode") in ("unavailable", "error"):
                with st.expander("Connection details (troubleshooting)"):
                    st.write(item.get("error") or "Local model unavailable.")
                    st.caption("Confirm Ollama is running and the configured model matches `ollama list`.")
            display_sources(item.get("sources", []), lang)
            last_user_text = next((m.get("content", "") for m in reversed(st.session_state.messages[:idx]) if m.get("role") == "user"), "")
            tts_lang = infer_voice_locale(last_user_text or item.get("content", ""), reply_lang)
            tts_button(item.get("content", ""), tts_lang, f"tts_v7_{idx}")
            st.markdown('<hr style="opacity:.40">', unsafe_allow_html=True)

    # Always-visible composer tools: no collapsed expander hiding document or microphone controls.
    st.markdown('<div class="chat-tools-heading"><span>＋</span><span>Attach a document or send a voice message</span><span class="muted">· optional</span></div>', unsafe_allow_html=True)
    doc_tab, voice_tab = st.tabs(["📎  Attach documents", "🎙  Voice message"])
    uploaded_files = []
    mic_audio = None
    uploaded_audio = None
    with doc_tab:
        uploaded_files = st.file_uploader(
            "Choose documents to discuss with SUDHAAR",
            type=["pdf", "docx", "txt", "md"], accept_multiple_files=True,
            key=f"document_uploads_v7_{st.session_state.upload_epoch}",
            help="Maximum 8 MB per file. Text is extracted locally; scanned-image PDFs are not OCR'd.",
        ) or []
        if uploaded_files:
            st.caption("Attached: " + " · ".join(f"{Path(f.name).name} ({f.size/1024:.0f} KB)" for f in uploaded_files))
        st.caption("Ask a question in the message box below, for example: ‘Summarise this document in simple Hindi.’ Uploaded text is unverified context, not medical evidence.")
    with voice_tab:
        st.write("Speak instead of typing. You can record directly or upload an existing voice note; SUDHAAR transcribes locally, then uses the transcript as your message.")
        mic_audio = None
        if hasattr(st, "audio_input"):
            mic_audio = st.audio_input("Record a voice message", sample_rate=16000, key=f"voice_input_v7_{st.session_state.audio_epoch}")
        else:
            st.info("Your Streamlit version does not include the microphone recorder. Upload an audio file below or update Streamlit in the project's virtual environment.")
        uploaded_audio = st.file_uploader(
            "Or upload a voice note", type=["wav", "mp3", "m4a", "mp4", "mpeg", "mpga", "webm", "ogg", "flac"],
            accept_multiple_files=False, key=f"voice_file_v7_{st.session_state.audio_epoch}",
            help="Local transcription requires the optional faster-whisper package.",
        )
        st.caption("Voice transcription setup: run `.venv\\Scripts\\python.exe -m pip install faster-whisper` in the VS Code terminal. The first use downloads a speech model.")

        def transcribe_for_chat(audio_bytes: bytes, suffix: str, send_now: bool) -> None:
            try:
                with st.spinner("Transcribing on this computer… the first use may download the speech model."):
                    transcript, detected_language = transcribe_audio(audio_bytes, suffix)
            except ImportError:
                st.error("Local voice transcription is optional. In the VS Code terminal run: .venv\\Scripts\\python.exe -m pip install faster-whisper")
                return
            except Exception as exc:
                st.error(f"Transcription failed ({type(exc).__name__}). You can type your message instead.")
                return

            if not transcript:
                st.warning("No clear speech was detected. Try a clearer recording or type your message.")
                return
            if send_now:
                upload_context, _, warnings = extract_uploaded_files(uploaded_files)
                st.session_state["voice_prompt_to_send_v7"] = transcript
                st.session_state["voice_context_to_send_v7"] = upload_context
                st.session_state["voice_warnings_to_show_v7"] = warnings
            else:
                st.session_state["pending_transcript"] = transcript
                st.session_state["transcribed_text"] = transcript
                st.session_state["transcript_language"] = detected_language or "unknown"
            # Keep rerun outside the try/except; rerun signals must not be treated as failures.
            st.rerun()

        mic_c1, mic_c2 = st.columns(2)
        with mic_c1:
            if st.button("Transcribe & review microphone", key=f"transcribe_mic_v7_{st.session_state.audio_epoch}", use_container_width=True, disabled=not bool(mic_audio)):
                transcribe_for_chat(mic_audio.getvalue(), ".wav", False)
        with mic_c2:
            if st.button("Transcribe & ask SUDHAAR", key=f"send_mic_v7_{st.session_state.audio_epoch}", use_container_width=True, type="primary", disabled=not bool(mic_audio)):
                transcribe_for_chat(mic_audio.getvalue(), ".wav", True)
        audio_c1, audio_c2 = st.columns(2)
        with audio_c1:
            if st.button("Transcribe & review audio file", key=f"transcribe_file_v7_{st.session_state.audio_epoch}", use_container_width=True, disabled=not bool(uploaded_audio)):
                transcribe_for_chat(uploaded_audio.getvalue(), Path(uploaded_audio.name).suffix.lower() or ".wav", False)
        with audio_c2:
            if st.button("Transcribe & ask SUDHAAR", key=f"send_file_v7_{st.session_state.audio_epoch}", use_container_width=True, type="primary", disabled=not bool(uploaded_audio)):
                transcribe_for_chat(uploaded_audio.getvalue(), Path(uploaded_audio.name).suffix.lower() or ".wav", True)

    # Handle the one-click voice flow outside the transcription try/except so a Streamlit rerun
    # cannot be caught and mislabeled as a transcription failure.
    queued_voice_prompt = st.session_state.pop("voice_prompt_to_send_v7", None)
    if queued_voice_prompt:
        queued_voice_context = st.session_state.pop("voice_context_to_send_v7", "")
        queued_voice_warnings = st.session_state.pop("voice_warnings_to_show_v7", [])
        for warning in queued_voice_warnings:
            st.warning(warning)
        process_user_message(
            queued_voice_prompt, reply_lang, lang,
            st.session_state.conversation_mode, st.session_state.response_style,
            queued_voice_context,
        )

    if st.session_state.get("pending_transcript"):
        st.markdown("**Review your voice transcript before sending**")
        st.caption(f"Detected language: {st.session_state.get('transcript_language', 'unknown')}. Correct recognition errors before sending.")
        st.text_area("Transcribed message", key="transcribed_text", height=100, max_chars=5000)
        tx1, tx2 = st.columns([3, 1])
        with tx1:
            if st.button("Send transcript to SUDHAAR ↗", key="send_transcript_v7", type="primary", use_container_width=True):
                transcript = st.session_state.get("transcribed_text", "").strip()
                upload_context, _, warnings = extract_uploaded_files(uploaded_files)
                for warning in warnings:
                    st.warning(warning)
                process_user_message(transcript, reply_lang, lang, st.session_state.conversation_mode, st.session_state.response_style, upload_context)
        with tx2:
            if st.button("Discard transcript", key="discard_transcript_v7", use_container_width=True):
                st.session_state["clear_transcript_next_run"] = True
                st.rerun()

    # Chat input remains pinned by Streamlit at the bottom, similar to common chat interfaces.
    typed_prompt = st.chat_input("Ask a question, type a topic like ‘alcohol’, or tell SUDHAAR what’s on your mind…", key="chat_composer_v7", max_chars=5000)
    queued_prompt = st.session_state.pop("queued_prompt_v6", None)
    prompt = typed_prompt or queued_prompt
    if prompt:
        upload_context, _, warnings = extract_uploaded_files(uploaded_files)
        for warning in warnings:
            st.warning(warning)
        process_user_message(prompt, reply_lang, lang, st.session_state.conversation_mode, st.session_state.response_style, upload_context)

    st.caption("SUDHAAR is a local-first educational prototype. It can make mistakes and does not diagnose or replace qualified care. General health answers and support conversations are different modes; use ‘Answer my question’ for direct information.")

elif page == "library":
    st.markdown('<div class="section-kicker">CURATED · SOURCE-LINKED · REVIEWABLE</div>', unsafe_allow_html=True)
    st.title("Health Library")
    st.write("Explore starter information, then open the original source. These entries are a working collection, not a complete or clinically validated medical reference.")
    query = st.text_input("Search topics or keywords", placeholder="Try: alcohol, stress, withdrawal, helpline, treatment", key="library_search")
    entries = retrieve(query, kb, top_k=12) if query.strip() else kb
    if not entries:
        st.info("No matching entry was found. Try a different word or open Help & Services.")
    for entry in entries:
        with st.container():
            st.markdown(f'<div class="kb-card scroll-reveal"><span class="pill">{safe_html_text(entry.get("topic", "General").replace("_", " ").upper())}</span><div class="card-title">{safe_html_text(entry.get("title", "Knowledge entry"))}</div><div class="card-text">{safe_html_text(entry.get("content", ""))}</div><p class="muted">{safe_html_text(entry.get("source_name", "Source not listed"))}</p></div>', unsafe_allow_html=True)
            if entry.get("source_url"):
                st.markdown(f"[Open original source ↗]({entry['source_url']})")
            with st.expander("Review metadata"):
                st.write(f"ID: {entry.get('id', 'N/A')} · Jurisdiction: {entry.get('jurisdiction', 'N/A')} · Last checked: {entry.get('last_verified', 'Not recorded')}")
                st.write(entry.get("review_status", "Review status not recorded"))

elif page == "services":
    st.markdown('<div class="section-kicker">HUMAN SUPPORT · OFFICIAL PATHWAYS</div>', unsafe_allow_html=True)
    st.title("Find support")
    st.write("SUDHAAR should connect people with real support, not try to replace it. Choose the pathway that fits your need.")
    st.markdown('<div class="hero" style="padding:22px 25px"><div class="eyebrow">IMMEDIATE DANGER</div><div class="call-number">112</div><p>If someone is unconscious, has trouble breathing, may have overdosed, or is in immediate life-threatening danger, contact emergency services now. Do not wait for an AI response.</p><a href="https://112.gov.in/">Official emergency information ↗</a></div>', unsafe_allow_html=True)
    search = st.text_input("Filter services", placeholder="mental health, de-addiction, reporting, government", key="service_search")
    filtered = [item for item in resources if not search or search.casefold() in " ".join(str(v) for v in item.values()).casefold()]
    cols = st.columns(2)
    for idx, item in enumerate(filtered):
        with cols[idx % 2]:
            number_html = f'<div class="call-number">{safe_html_text(item.get("number", ""))}</div>' if item.get("number") else ""
            card_html = (f'<div class="resource-card scroll-reveal"><span class="pill">{safe_html_text(item.get("kind", "Support"))}</span>'
                         f'<div class="card-title">{safe_html_text(item.get("name", "Service"))}</div>{number_html}'
                         f'<div class="card-text">{safe_html_text(item.get("description", ""))}</div>'
                         f'<p class="muted">{safe_html_text(item.get("source_name", "Official source"))} · Last checked: '
                         f'{safe_html_text(item.get("checked_on", item.get("verified_on", "Not recorded")))}</p></div>')
            st.markdown(card_html, unsafe_allow_html=True)
            actions = st.columns(2)
            with actions[0]:
                if item.get("number") and item.get("phone_link"):
                    st.link_button(f"☎ Call {item['number']}", item["phone_link"], use_container_width=True)
            with actions[1]:
                st.link_button("Open official page ↗", item.get("website", "https://www.india.gov.in/"), use_container_width=True)
    st.caption("14446: de-addiction information and referral. 14416: Tele-MANAS mental-health support. 1933: narcotics-related reporting, not counselling. Reverify numbers and scope before public use.")

elif page == "programmes":
    st.markdown('<div class="section-kicker">CONNECT, DON’T DUPLICATE</div>', unsafe_allow_html=True)
    st.title("Programmes & digital services")
    st.write("Official links that can help users reach established services. SUDHAAR is a bridge, not a replacement for these programmes.")
    programs = [
        {"title": "Nasha Mukt Bharat Abhiyaan", "type": "Substance-use prevention & support", "desc": "Official programme information and public-awareness resources.", "url": "https://nmba.dosje.gov.in/", "badge": "Substance use"},
        {"title": "National De-addiction Helpline · 14446", "type": "Counselling / referral", "desc": "Official information about support for people affected by substance use and their families.", "url": "https://socialjustice.gov.in/schemes/42", "badge": "14446 support"},
        {"title": "Tele-MANAS · 14416", "type": "Mental-health support", "desc": "Government tele-mental-health service and referral information.", "url": "https://telemanas.mohfw.gov.in/", "badge": "Mental health"},
        {"title": "eSanjeevani", "type": "Telemedicine", "desc": "Official telemedicine platform. Check eligibility and available services; not an emergency channel.", "url": "https://esanjeevani.mohfw.gov.in/", "badge": "Telemedicine"},
        {"title": "Ayushman Bharat Digital Mission", "type": "Digital-health ecosystem", "desc": "Official information about India's digital-health ecosystem; not addiction-specific treatment.", "url": "https://abdm.gov.in/", "badge": "Digital health"},
        {"title": "Emergency Response Support System · 112", "type": "Emergency response", "desc": "Official information on India's unified emergency number.", "url": "https://112.gov.in/", "badge": "Emergency"},
        {"title": "MANAS National Narcotics Helpline · 1933", "type": "Drug-related reporting", "desc": "For drug-related reporting, including illegal trafficking or sale; not a counselling service.", "url": "https://www.india.gov.in/category/welfare-of-families/subcategory/family-welfare/details/website-of-manas-national-narcotics-helpline", "badge": "Reporting"},
        {"title": "NIMHANS", "type": "Specialist institution", "desc": "Official mental-health and neuroscience institution website for research, education and service information.", "url": "https://nimhans.ac.in/", "badge": "Specialist"},
        {"title": "Ministry of Health & Family Welfare", "type": "Public-health portal", "desc": "Official ministry portal and public-health programme updates.", "url": "https://www.mohfw.gov.in/", "badge": "Public health"},
        {"title": "World Health Organization · mhGAP", "type": "Evidence and guidelines", "desc": "WHO guideline for mental, neurological and substance-use conditions.", "url": "https://www.who.int/publications/i/item/9789240084278", "badge": "Guidelines"},
    ]
    filter_text = st.text_input("Search programmes", placeholder="Try: counselling, emergency, digital health, treatment", key="program_search")
    shown = [p for p in programs if not filter_text or filter_text.casefold() in " ".join(p.values()).casefold()]
    cols = st.columns(2)
    for idx, item in enumerate(shown):
        with cols[idx % 2]:
            st.markdown(f'<div class="programme-card scroll-reveal"><span class="pill">{safe_html_text(item["badge"])}</span><div class="card-title">{safe_html_text(item["title"])}</div><div class="muted">{safe_html_text(item["type"])}</div><div class="card-text">{safe_html_text(item["desc"])}</div></div>', unsafe_allow_html=True)
            st.link_button("Open official resource ↗", item["url"], use_container_width=True)

elif page == "plan":
    st.markdown('<div class="section-kicker">OPTIONAL · USER-CONTROLLED · SESSION ONLY</div>', unsafe_allow_html=True)
    st.title("My next step")
    st.write("You do not have to solve everything at once. Choose one small action that feels possible for you today.")
    st.info("This is a non-clinical support plan, not a treatment or detox plan. It stays only in the current app session.")
    with st.form("support_plan_form"):
        goal = st.selectbox("What would feel useful today?", ["Take one manageable next step", "Understand treatment options", "Find a professional support service", "Learn about health effects", "Support a family member", "Plan what to do when overwhelmed", "Understand government programmes"], key="plan_goal_widget")
        supporter = st.selectbox("What kind of support feels appropriate?", ["A trusted person or counsellor", "A qualified healthcare professional", "A government helpline", "I am not sure yet"], key="plan_support_widget")
        steps = st.multiselect("Pick possible next steps", ["Read a trusted health resource", "Contact a qualified professional", "Call 14446 for de-addiction support", "Call 14416 for mental-health support", "Talk with someone I trust", "Open an official government service"], key="plan_steps_widget")
        save = st.form_submit_button("Save my session-only plan", use_container_width=True)
    if save:
        st.session_state.support_goal = goal
        st.session_state.support_person = supporter
        st.session_state.support_steps = steps
        st.session_state.support_plan_ready = True
    if st.session_state.support_plan_ready:
        steps_html = "".join(f'<li>{safe_html_text(step)}</li>' for step in st.session_state.support_steps)
        if not steps_html:
            steps_html = '<li>No next steps selected yet. You can choose later.</li>'
        plan_html = (f'<div class="plan-card"><div class="section-kicker">YOUR PLAN</div>'
                     f'<div class="card-title">{safe_html_text(st.session_state.support_goal)}</div>'
                     f'<div class="card-text">Preferred support: {safe_html_text(st.session_state.support_person)}</div>'
                     f'<ul>{steps_html}</ul></div>')
    else:
        plan_html = '<div class="plan-card"><div class="section-kicker">YOUR PLAN</div><div class="card-title">There is no perfect next step.</div><div class="card-text">Start with something small and safe. Choose options above and save when you are ready.</div></div>'
    st.markdown(plan_html, unsafe_allow_html=True)
    st.caption("For immediate danger or serious symptoms, contact emergency services; do not rely on this plan.")

elif page == "settings":
    st.markdown('<div class="section-kicker">MAKE THIS SPACE FEEL RIGHT FOR YOU</div>', unsafe_allow_html=True)
    st.title("Preferences")
    st.write("Change the appearance or response style any time. Your place in the app and your current session conversation should stay where you left them.")
    a, b = st.columns(2)
    with a:
        st.markdown('<div class="feature-card"><div class="icon-box">◐</div><div class="card-title">Appearance</div><div class="card-text">Midnight is a low-light dark theme. Pearl is a high-contrast light theme.</div></div>', unsafe_allow_html=True)
        st.markdown(f"**Current theme:** {'🌙 Midnight' if st.session_state.theme_mode == 'Midnight' else '☀️ Pearl'}")
        st.caption("Use the Appearance selector in the sidebar to switch themes. Your current page and chat stay selected.")
        st.selectbox("Preferred answer style", ["Gentle and reassuring", "Clear and concise", "Detailed explanation"], key="response_style")
    with b:
        st.markdown('<div class="feature-card"><div class="icon-box">🔒</div><div class="card-title">Privacy by default</div><div class="card-text">This prototype uses in-memory session state and does not intentionally write chats or uploaded text into a database.</div></div>', unsafe_allow_html=True)
        st.write("The local model processes questions through Ollama on your machine. Voice transcription, when enabled with faster-whisper, also runs locally. Browser text-to-speech uses voices installed in your browser/operating system.")
        st.write("Do not upload identifying health records or details you do not want processed. This prototype has not undergone a formal privacy/security audit.")
    st.markdown('<div class="feature-card"><div class="card-title">Language quality</div><div class="card-text">AI replies can be requested in the 22 scheduled Indian languages plus English. The interface has partial translations in 13 languages. Medical and crisis translations need fluent-speaker review; availability in the selector does not guarantee accuracy.</div></div>', unsafe_allow_html=True)
    st.markdown("**Local AI status**")
    status = cached_ollama_status()
    st.json({"server_running": status["server"], "model_installed": status["model"], "model": status["model_name"], "endpoint": os.getenv("OLLAMA_BASE_URL", "http://127.0.0.1:11434")})
    st.button("Clear current conversation", key="settings_clear_chat", on_click=clear_chat_and_go_to_chat, use_container_width=True)

elif page == "feedback":
    st.markdown('<div class="section-kicker">HELP US MAKE THIS SAFER AND MORE USEFUL</div>', unsafe_allow_html=True)
    st.title("Feedback")
    st.write("Tell us what would make the experience kinder, clearer or easier to use. Don't include any personal health information.")
    with st.form("feedback_form_v4"):
        area = st.selectbox("What should we improve?", ["UI and visual design", "Compassion / conversation tone", "Language and translation", "Knowledge quality and sources", "AI answer quality", "Voice input or read-aloud", "Safety and referrals", "Government-service directory", "Other"])
        rating = st.select_slider("How easy was this prototype to use?", options=["Difficult", "Okay", "Easy", "Very easy"])
        note = st.text_area("Suggestion (no personal/medical details)", max_chars=800)
        submitted = st.form_submit_button("Send feedback", use_container_width=True)
    if submitted:
        st.session_state.feedback_done = True
        st.success("Thank you. Feedback is only acknowledged in this running session; this prototype does not send or persist it.")
        st.caption(f"Session-only summary: {area} · usability={rating} · {len(note)} suggestion characters.")

st.markdown("<hr>", unsafe_allow_html=True)
st.caption("SUDHAAR · local-first prototype · supportive education, not a diagnosis, prescription, emergency response or treatment service. Verify source and contact details before public deployment.")
